"""Single revised slide: 'How the encoder learned: fill in the blank'.

Standalone draft to iterate on (-> slides/encoder_revised.pptx). BERT-style:
colored element tokens, one MASKed, the encoder (being trained, not frozen)
predicts the hidden amount, loss compared to truth, backprop loop to the weights.

    conda run -n mineral python slides/encoder_revised.py
"""

from pathlib import Path

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

INK = "1A2238"
BLUE = "2E5BFF"
TEAL = "12B5A5"
PURPLE = "6C5CE7"
ORANGE = "E8883A"
RED = "C0392B"
GRAYL = "C8CDD6"
GRAYT = "8A93A3"
WHITE = "FFFFFF"
MASKG = "9AA3B2"

W, H = Inches(13.333), Inches(7.5)


def _rgb(h):
    return RGBColor.from_string(h)


def _lighten(h, amt=0.85):
    r, g, b = (int(h[i:i + 2], 16) for i in (0, 2, 4))
    r, g, b = (int(c + (255 - c) * amt) for c in (r, g, b))
    return f"{r:02X}{g:02X}{b:02X}"


def _put(tf, lines, size, color, bold=True, align=PP_ALIGN.CENTER):
    tf.word_wrap = True
    for i, ln in enumerate(lines if isinstance(lines, list) else [lines]):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        sub = isinstance(ln, tuple)
        run = p.add_run()
        run.text = ln[0] if sub else ln
        run.font.size = Pt(size - 5 if sub else size)
        run.font.bold = bold and not sub
        run.font.color.rgb = _rgb(GRAYT if sub else color)


def box(s, x, y, w, h, lines, color, fill=None, size=13, rounded=True):
    shape = MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE
    shp = s.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    shp.fill.solid()
    shp.fill.fore_color.rgb = _rgb(fill if fill else _lighten(color, 0.86))
    shp.line.color.rgb = _rgb(color)
    shp.line.width = Pt(2)
    shp.shadow.inherit = False
    shp.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    _put(shp.text_frame, lines, size, INK)
    return shp


def token(s, x, y, w, h, sym, amt, color, masked=False):
    shp = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    shp.fill.solid()
    shp.fill.fore_color.rgb = _rgb(MASKG if masked else _lighten(color, 0.78))
    shp.line.color.rgb = _rgb(MASKG if masked else color)
    shp.line.width = Pt(2.5 if masked else 1.75)
    shp.line.dash_style = 2 if masked else None
    shp.shadow.inherit = False
    tf = shp.text_frame
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    if masked:
        _put(tf, ["[MASK]", ("hide amount",)], 17, WHITE)
    else:
        _put(tf, [sym, (amt,)], 19, INK)
    return shp


def arrow(s, x, y, w, h, color=INK, shape=MSO_SHAPE.DOWN_ARROW):
    a = s.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    a.fill.solid(); a.fill.fore_color.rgb = _rgb(color)
    a.line.fill.background(); a.shadow.inherit = False
    return a


def line(s, x1, y1, x2, y2, color=RED, width=2.0, dash=None):
    c = s.shapes.add_connector(2, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    c.line.color.rgb = _rgb(color); c.line.width = Pt(width)
    if dash:
        c.line.dash_style = dash
    c.shadow.inherit = False
    return c


def txt(s, x, y, w, lines, size, color, align=PP_ALIGN.CENTER, bold=False):
    t = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(0.6))
    _put(t.text_frame, lines, size, color, bold=bold, align=align)


def build():
    prs = Presentation()
    prs.slide_width, prs.slide_height = W, H
    s = prs.slides.add_slide(prs.slide_layouts[6])

    txt(s, 0.55, 0.3, 12, "ACT 2  -  B4", 13, BLUE, PP_ALIGN.LEFT, bold=True)
    txt(s, 0.55, 0.66, 12.2, "How the encoder learned: fill in the blank", 26, INK,
        PP_ALIGN.LEFT, bold=True)
    txt(s, 0.55, 1.18, 12.2, "like predicting a missing word from its sentence -- "
        "but the 'words' are elements and the blank is an amount", 13, GRAYT, PP_ALIGN.LEFT)

    # token row: 6 majors-ish + 1 trace; mask the 4th
    toks = [("Si", "0.50", BLUE), ("Al", "0.16", BLUE), ("Ca", "0.10", BLUE),
            ("Mg", "0.075", BLUE), ("Fe", "0.11", BLUE), ("Na", "0.025", BLUE),
            ("Sr", "4e-4", TEAL)]
    masked_idx = 3
    tw, th, ty = 1.5, 1.15, 1.95
    x0 = (13.333 - (len(toks) * tw + (len(toks) - 1) * 0.25)) / 2
    xs = [x0 + i * (tw + 0.25) for i in range(len(toks))]
    for i, (sym, amt, col) in enumerate(toks):
        token(s, xs[i], ty, tw, th, sym, amt, col, masked=(i == masked_idx))
    txt(s, xs[0] - 1.4, ty + 0.35, 1.3, ["a rock =", ("element tokens",)], 12, GRAYT, PP_ALIGN.RIGHT)

    # tokens -> encoder
    enc_y, enc_h = 3.85, 1.0
    for i in range(len(toks)):
        cx = xs[i] + tw / 2
        line(s, cx, ty + th, cx, enc_y, color=(MASKG if i == masked_idx else GRAYL), width=1.5)
    box(s, x0, enc_y, len(toks) * tw + (len(toks) - 1) * 0.25, enc_h,
        ["Transformer encoder", ("being trained -- weights updated each step (not frozen)",)],
        PURPLE, size=16)

    # encoder -> readout of the masked slot
    mcx = xs[masked_idx] + tw / 2
    arrow(s, mcx - 0.16, enc_y + enc_h + 0.05, 0.32, 0.4, color=PURPLE)
    py = 5.5
    box(s, 2.4, py, 3.0, 1.05, ["encoder predicts", ("Mg amount ~ 0.071",)], ORANGE, size=14)
    box(s, 5.7, py, 1.9, 1.05, ["truth", ("0.075",)], GRAYT, fill=WHITE, size=14)
    arrow(s, 5.45, py + 0.32, 0.3, 0.4, color=INK, shape=MSO_SHAPE.RIGHT_ARROW)
    arrow(s, 7.65, py + 0.32, 0.3, 0.4, color=INK, shape=MSO_SHAPE.RIGHT_ARROW)
    box(s, 8.0, py, 3.0, 1.05, ["log-space", "MSE loss"], RED, size=15)

    # backprop loop: loss -> up the right side -> into encoder, weights updated
    line(s, 11.0, py + 0.5, 12.65, py + 0.5, RED, 2.5)
    line(s, 12.65, py + 0.5, 12.65, enc_y + enc_h / 2, RED, 2.5)
    arrow(s, 12.1, enc_y + enc_h / 2 - 0.16, 0.55, 0.32, color=RED, shape=MSO_SHAPE.LEFT_ARROW)
    txt(s, 9.7, 4.55, 3.4, ["backpropagation:", "updates encoder weights"], 12, RED, PP_ALIGN.CENTER, bold=True)

    txt(s, 0.55, 6.75, 12.2, "mask one element's amount  ->  predict it from the rest  ->  "
        "the error backpropagates and adjusts the encoder.  Repeat over ~167k rocks.",
        13, INK, PP_ALIGN.LEFT)

    out = Path(__file__).resolve().parent / "encoder_revised.pptx"
    prs.save(out)
    print(f"wrote {out}")


if __name__ == "__main__":
    build()
